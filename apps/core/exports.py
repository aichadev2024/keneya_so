"""
Génération d'exports PDF/Excel génériques, réutilisables par tous les modules
(CDC 2.2 « rapports exportables (PDF/Excel) » / CDC 9 « export PDF/Excel
possible sur tous les modules »).

Choix technique — **ReportLab** pour le PDF plutôt que WeasyPrint : WeasyPrint
nécessite des bibliothèques système (Pango, Cairo, GDK-Pixbuf) absentes de
cette machine de développement et dont l'installation requiert des droits
d'administration indisponibles ici (même contrainte que pour GNU gettext,
voir README « Internationalisation »). ReportLab est pur Python — aucune
dépendance système — et figure explicitement parmi les choix du CDC (section 8).
Excel : **openpyxl**, également pur Python.

Rendu de l'arabe en PDF : ReportLab ne fait pas la liaison des lettres arabes
ni l'inversion visuelle de droite à gauche par défaut ; ``arabic_reshaper`` +
``python-bidi`` corrigent le texte avant impression, mais encore faut-il une
police couvrant les glyphes arabes. Ce module tente d'enregistrer une police
Unicode déjà présente sur la machine (ex. Tahoma/Segoe UI sous Windows,
DejaVu/Noto sous Linux) et se replie sur Helvetica (latin uniquement) si
aucune n'est trouvée. **Pour une mise en production servant des rapports en
arabe, installer une police libre couvrant l'arabe (ex. paquet Debian/Ubuntu
``fonts-noto-core``) sur le serveur.** Les exports Excel n'ont pas cette
limite : openpyxl stocke du texte Unicode brut, rendu par la police du poste
qui ouvre le fichier.
"""

from __future__ import annotations

import os
import re
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Font as XlFont, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COULEUR_ENTETE = colors.HexColor("#1F3864")  # cohérent avec static/css/app.css

_POLICES_UNICODE = [
    ("Tahoma", r"C:\Windows\Fonts\tahoma.ttf"),
    ("SegoeUI", r"C:\Windows\Fonts\segoeui.ttf"),
    ("ArialSysteme", r"C:\Windows\Fonts\arial.ttf"),
    ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ("NotoSansArabic", "/usr/share/fonts/opentype/noto/NotoSansArabic-Regular.ttf"),
    ("NotoSans", "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"),
]
_ARABE = re.compile(r"[\u0600-\u06FF]")
_police_cache: str | None = None


def _police_pdf() -> str:
    """Enregistre et renvoie le nom d'une police Unicode disponible (best-effort)."""
    global _police_cache
    if _police_cache:
        return _police_cache
    for nom, chemin in _POLICES_UNICODE:
        if os.path.exists(chemin):
            try:
                pdfmetrics.registerFont(TTFont(nom, chemin))
                _police_cache = nom
                return nom
            except Exception:
                continue
    _police_cache = "Helvetica"
    return _police_cache


def _texte(valeur) -> str:
    """Convertit une valeur de cellule en texte, avec mise en forme RTL pour l'arabe."""
    if valeur is None:
        texte = "—"
    elif isinstance(valeur, bool):
        texte = "✓" if valeur else "—"
    elif isinstance(valeur, datetime):
        texte = timezone.localtime(valeur).strftime("%d/%m/%Y %H:%M")
    elif isinstance(valeur, date):
        texte = valeur.strftime("%d/%m/%Y")
    else:
        texte = str(valeur)
    if _ARABE.search(texte):
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display
            return get_display(arabic_reshaper.reshape(texte))
        except Exception:
            return texte
    return texte


_TYPES_EXCEL_NATIFS = (str, int, float, bool, date, datetime)


def _valeur_excel(v):
    """Convertit une valeur de cellule vers un type qu'openpyxl accepte nativement."""
    if v is None:
        return ""
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return timezone.localtime(v).replace(tzinfo=None) if timezone.is_aware(v) else v
    if isinstance(v, _TYPES_EXCEL_NATIFS):
        return v
    return str(v)  # ex. : chaînes de traduction paresseuses (gettext_lazy)


def _nom_fichier(base: str, extension: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]+", "-", base).strip("-") or "export"
    horodatage = timezone.localtime().strftime("%Y%m%d-%H%M")
    return f"{safe}-{horodatage}.{extension}"


def exporter_pdf(*, titre: str, colonnes: list[str], lignes: list[list],
                 sous_titre: str = "", nom_fichier: str = "export",
                 paysage: bool = True) -> HttpResponse:
    """
    Construit un PDF tabulaire générique (utilisable par n'importe quel module).

    ``lignes`` : liste de séquences de valeurs brutes (str, nombre, date, bool…),
    dans l'ordre des ``colonnes``.
    """
    import io

    tampon = io.BytesIO()
    taille_page = landscape(A4) if paysage else A4
    doc = SimpleDocTemplate(
        tampon, pagesize=taille_page,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title=titre,
    )
    police = _police_pdf()
    styles = getSampleStyleSheet()
    style_titre = ParagraphStyle("TitreKS", parent=styles["Title"], fontName=police,
                                 textColor=COULEUR_ENTETE, fontSize=16)
    style_sous_titre = ParagraphStyle("SousTitreKS", parent=styles["Normal"],
                                      fontName=police, fontSize=9,
                                      textColor=colors.grey)

    elements = [
        Paragraph(_texte(titre), style_titre),
        Paragraph(_texte(sous_titre or
                         f"Généré le {timezone.localtime():%d/%m/%Y %H:%M} — Kènèya Sô"),
                 style_sous_titre),
        Spacer(1, 0.5 * cm),
    ]

    donnees = [[_texte(c) for c in colonnes]]
    donnees += [[_texte(v) for v in ligne] for ligne in lignes]

    largeur_dispo = taille_page[0] - 3 * cm
    largeur_colonne = largeur_dispo / max(len(colonnes), 1)
    table = Table(donnees, colWidths=[largeur_colonne] * len(colonnes), repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), police),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), COULEUR_ENTETE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), police),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)
    if len(lignes) == 0:
        elements.append(Spacer(1, 0.5 * cm))
        elements.append(Paragraph(_texte("Aucune donnée pour cette sélection."),
                                  styles["Normal"]))

    doc.build(elements)
    tampon.seek(0)
    reponse = HttpResponse(tampon.read(), content_type="application/pdf")
    reponse["Content-Disposition"] = f'inline; filename="{_nom_fichier(nom_fichier, "pdf")}"'
    return reponse


def exporter_pdf_multi(*, titre: str, sections: list[tuple[str, list[str], list[list]]],
                       sous_titre: str = "", nom_fichier: str = "export") -> HttpResponse:
    """
    Variante multi-tableaux d'``exporter_pdf`` (rapports composites : indicateurs
    du bloc opératoire — CDC 5.3.7 — ou tout autre tableau de bord).

    ``sections`` : liste de ``(titre_section, colonnes, lignes)``.
    """
    import io

    tampon = io.BytesIO()
    taille_page = landscape(A4)
    doc = SimpleDocTemplate(
        tampon, pagesize=taille_page,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title=titre,
    )
    police = _police_pdf()
    styles = getSampleStyleSheet()
    style_titre = ParagraphStyle("TitreKS", parent=styles["Title"], fontName=police,
                                 textColor=COULEUR_ENTETE, fontSize=16)
    style_sous_titre = ParagraphStyle("SousTitreKS", parent=styles["Normal"],
                                      fontName=police, fontSize=9, textColor=colors.grey)
    style_section = ParagraphStyle("SectionKS", parent=styles["Heading2"], fontName=police,
                                   textColor=COULEUR_ENTETE, fontSize=12,
                                   spaceBefore=10, spaceAfter=4)

    elements = [
        Paragraph(_texte(titre), style_titre),
        Paragraph(_texte(sous_titre or
                         f"Généré le {timezone.localtime():%d/%m/%Y %H:%M} — Kènèya Sô"),
                 style_sous_titre),
    ]
    largeur_dispo = taille_page[0] - 3 * cm

    for titre_section, colonnes, lignes in sections:
        elements.append(Paragraph(_texte(titre_section), style_section))
        if not lignes:
            elements.append(Paragraph(_texte("Aucune donnée."), styles["Normal"]))
            continue
        donnees = [[_texte(c) for c in colonnes]] + [[_texte(v) for v in l] for l in lignes]
        largeur_colonne = largeur_dispo / max(len(colonnes), 1)
        table = Table(donnees, colWidths=[largeur_colonne] * len(colonnes), repeatRows=1)
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), police),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), COULEUR_ENTETE),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(table)

    doc.build(elements)
    tampon.seek(0)
    reponse = HttpResponse(tampon.read(), content_type="application/pdf")
    reponse["Content-Disposition"] = f'inline; filename="{_nom_fichier(nom_fichier, "pdf")}"'
    return reponse


def exporter_excel_multi(*, sections: list[tuple[str, list[str], list[list]]],
                         nom_fichier: str = "export") -> HttpResponse:
    """Variante multi-feuilles d'``exporter_excel`` (une feuille par section)."""
    classeur = openpyxl.Workbook()
    classeur.remove(classeur.active)

    noms_utilises: set[str] = set()
    for titre_section, colonnes, lignes in sections:
        nom_feuille = re.sub(r"[:\\/?*\[\]]", "-", str(titre_section) or "Feuille")[:31]
        base, suffixe = nom_feuille, 2
        while nom_feuille in noms_utilises:
            nom_feuille = f"{base[:28]}-{suffixe}"
            suffixe += 1
        noms_utilises.add(nom_feuille)

        feuille = classeur.create_sheet(title=nom_feuille)
        feuille.append([_valeur_excel(c) for c in colonnes])
        for cellule in feuille[1]:
            cellule.font = XlFont(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor="1F3864")
        feuille.freeze_panes = "A2"
        for ligne in lignes:
            feuille.append([_valeur_excel(v) for v in ligne])
        for i, colonne in enumerate(colonnes, start=1):
            lettre = get_column_letter(i)
            largeur = max(len(str(colonne)), 10)
            for cellule in feuille[lettre][1:200]:
                if cellule.value is not None:
                    largeur = max(largeur, len(str(cellule.value)))
            feuille.column_dimensions[lettre].width = min(largeur + 2, 60)

    import io
    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    reponse = HttpResponse(
        tampon.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    reponse["Content-Disposition"] = f'attachment; filename="{_nom_fichier(nom_fichier, "xlsx")}"'
    return reponse


def exporter_excel(*, titre: str, colonnes: list[str], lignes: list[list],
                   nom_fichier: str = "export") -> HttpResponse:
    """Construit un classeur Excel générique (une feuille, en-tête figé)."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = re.sub(r"[:\\/?*\[\]]", "-", str(titre))[:31] or "Export"

    feuille.append([_valeur_excel(c) for c in colonnes])
    for cellule in feuille[1]:
        cellule.font = XlFont(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F3864")
        cellule.alignment = Alignment(vertical="center")
    feuille.freeze_panes = "A2"

    for ligne in lignes:
        feuille.append([_valeur_excel(v) for v in ligne])

    for i, colonne in enumerate(colonnes, start=1):
        lettre = get_column_letter(i)
        largeur = max(len(str(colonne)), 10)
        for cellule in feuille[lettre][1:200]:
            if cellule.value is not None:
                largeur = max(largeur, len(str(cellule.value)))
        feuille.column_dimensions[lettre].width = min(largeur + 2, 60)

    import io
    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    reponse = HttpResponse(
        tampon.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    reponse["Content-Disposition"] = f'attachment; filename="{_nom_fichier(nom_fichier, "xlsx")}"'
    return reponse


def exporter(request, *, titre, colonnes, lignes, sous_titre="", nom_fichier="export"):
    """
    Sélectionne le format d'export depuis ``?format=pdf|xlsx`` et délègue.
    Renvoie ``None`` si aucun format d'export n'est demandé (affichage HTML normal).
    """
    fmt = request.GET.get("format")
    if fmt == "pdf":
        return exporter_pdf(titre=titre, colonnes=colonnes, lignes=lignes,
                            sous_titre=sous_titre, nom_fichier=nom_fichier)
    if fmt in {"xlsx", "excel"}:
        return exporter_excel(titre=titre, colonnes=colonnes, lignes=lignes,
                              nom_fichier=nom_fichier)
    return None


def exporter_multi(request, *, titre, sections, sous_titre="", nom_fichier="export"):
    """Variante multi-tableaux d'``exporter`` (voir ``exporter_pdf_multi``)."""
    fmt = request.GET.get("format")
    if fmt == "pdf":
        return exporter_pdf_multi(titre=titre, sections=sections, sous_titre=sous_titre,
                                  nom_fichier=nom_fichier)
    if fmt in {"xlsx", "excel"}:
        return exporter_excel_multi(sections=sections, nom_fichier=nom_fichier)
    return None


class ExportableListMixin:
    """
    Ajoute l'export PDF/Excel générique (``?format=pdf`` / ``?format=xlsx``) à
    une ``ListView`` existante, sur le même queryset filtré que l'affichage
    HTML (CDC 9 — export possible sur tous les modules).

    Les sous-classes définissent ``export_titre``, ``export_nom_fichier``,
    ``export_colonnes()`` et ``export_ligne(objet)``.
    """

    export_titre = "Export"
    export_nom_fichier = "export"

    def export_colonnes(self) -> list[str]:
        raise NotImplementedError

    def export_ligne(self, objet) -> list:
        raise NotImplementedError

    def get(self, request, *args, **kwargs):
        if request.GET.get("format") in {"pdf", "xlsx", "excel"}:
            self.object_list = self.get_queryset()
            lignes = [self.export_ligne(o) for o in self.object_list]
            reponse = exporter(
                request, titre=str(self.export_titre),
                colonnes=self.export_colonnes(), lignes=lignes,
                nom_fichier=self.export_nom_fichier,
            )
            if reponse is not None:
                return reponse
        return super().get(request, *args, **kwargs)
