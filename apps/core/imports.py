"""
Import Excel générique, symétrique à ``exports.py`` — permet à un hôpital qui a
déjà sa liste de médicaments ou d'examens (dans un tableur) de l'importer en
masse au lieu de ressaisir chaque fiche une à une.

Choix technique : **openpyxl** uniquement (déjà une dépendance du projet, pur
Python). Pas de lecture PDF/Word ici : la mise en page de ces documents varie
trop d'un hôpital à l'autre pour une extraction fiable sans réglage par cas —
voir la discussion dans le README / les échanges avec la cliente. On fournit
à la place un modèle Excel téléchargeable, facile à remplir en copiant-collant
depuis un document existant.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from django.http import HttpResponse

import openpyxl
from openpyxl.styles import Font as XlFont, PatternFill
from openpyxl.utils import get_column_letter

_ENTETE_COULEUR = "1F3864"  # cohérent avec exports.py / static/css/app.css


@dataclass
class ResultatImport:
    """Bilan d'un import : nombre de lignes créées/ignorées et messages d'erreur."""

    crees: int = 0
    ignores: int = 0
    erreurs: list[str] = field(default_factory=list)

    @property
    def total_erreurs(self) -> int:
        return len(self.erreurs)


def generer_modele_excel(*, colonnes: list[str], exemple: list | None = None,
                         nom_fichier: str = "modele") -> HttpResponse:
    """Classeur Excel vierge avec les bons en-têtes (+ une ligne d'exemple facultative)."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "Import"
    feuille.append(colonnes)
    for cellule in feuille[1]:
        cellule.font = XlFont(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor=_ENTETE_COULEUR)
    feuille.freeze_panes = "A2"
    if exemple:
        feuille.append(exemple)
    for i, colonne in enumerate(colonnes, start=1):
        feuille.column_dimensions[get_column_letter(i)].width = max(len(str(colonne)) + 4, 14)

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    reponse = HttpResponse(
        tampon.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier}.xlsx"'
    return reponse


def lire_lignes_excel(fichier) -> tuple[list[str], list[list]]:
    """
    Lit un classeur ``.xlsx`` uploadé (première feuille) : la première ligne est
    traitée comme l'en-tête, les lignes entièrement vides sont ignorées.
    Renvoie ``(en_tetes, lignes)``.
    """
    classeur = openpyxl.load_workbook(fichier, data_only=True, read_only=True)
    feuille = classeur.worksheets[0]
    lignes_brutes = list(feuille.iter_rows(values_only=True))
    if not lignes_brutes:
        return [], []
    en_tetes = [str(c).strip() if c is not None else "" for c in lignes_brutes[0]]
    lignes = [list(l) for l in lignes_brutes[1:] if any(c not in (None, "") for c in l)]
    return en_tetes, lignes


def valeur_texte(v) -> str:
    """Normalise une cellule en texte (gère les nombres/flottants venus d'Excel)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def valeur_booleenne(v, defaut: bool = False) -> bool:
    """Interprète « oui/non », « vrai/faux », 1/0, True/False… venus d'une cellule Excel."""
    texte = valeur_texte(v).strip().lower()
    if not texte:
        return defaut
    return texte in {"oui", "o", "vrai", "true", "1", "x", "yes", "y"}


def valeur_entiere(v, defaut: int) -> int:
    texte = valeur_texte(v)
    if not texte:
        return defaut
    try:
        return int(float(texte))
    except ValueError:
        return defaut
